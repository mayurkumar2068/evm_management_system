"""Module B: Dashboard and reports (5 APIs). Scope comes from the token."""
import csv
import io
from datetime import datetime
from typing import Literal

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from sqlalchemy import and_, case, func, select
from sqlalchemy.orm import Session

from app.core.deps import OFFICIAL_ROLES, require_roles
from app.core.errors import forbidden
from app.core.timeutil import start_of_today_utc, as_utc, utcnow
from app.db.session import get_db
from app.models import (Assignment, AuditLog, Complaint, ComplaintStageSla, Constituency, Decision, District,
                        FlyingSquad, User, ViolationType)
from app.services import sla
from app.services.access import assert_can_view, get_complaint

router = APIRouter(tags=["B. Dashboard & reports"])
dash_roles = require_roles("DC", "RO", "DEO", "CEO", "ADMIN")

OPEN = ("RECEIVED", "ASSIGNED", "ACCEPTED", "EN_ROUTE", "REACHED", "REPORT_SUBMITTED", "REPORT_RETURNED", "ESCALATED")


def _scope_conditions(db: Session, user: User, district: str | None, ac: str | None) -> list:
    conds = []
    if user.role in ("DC", "DEO"):
        conds.append(Complaint.district_id == (user.district_id or -1))
    elif user.role == "RO":
        conds.append(Complaint.ac_id == (user.ac_id or -1))
    if district:
        if user.role not in ("CEO", "ADMIN"):
            own = db.scalar(select(District.code).where(District.id == user.district_id))
            if own != district:
                raise forbidden("District outside your scope")
        conds.append(Complaint.district_id == select(District.id).where(District.code == district).scalar_subquery())
    if ac:
        conds.append(Complaint.ac_id == select(Constituency.id).where(Constituency.code == ac).scalar_subquery())
    return conds


def _range(date_from: datetime | None, date_to: datetime | None) -> tuple[datetime, datetime]:
    return (as_utc(date_from) if date_from else start_of_today_utc(),
            as_utc(date_to) if date_to else utcnow())


@router.get("/dashboard/summary", summary="B1 KPI tiles")
def summary(district: str | None = None, ac: str | None = None,
            user: User = Depends(dash_roles), db: Session = Depends(get_db)):
    conds = _scope_conditions(db, user, district, ac)
    today = start_of_today_utc()

    def count(*extra):
        return db.scalar(select(func.count()).select_from(Complaint).where(*conds, *extra))

    return {"as_of": utcnow(),
            "received_today": count(Complaint.submitted_at >= today),
            "awaiting_allocation": count(Complaint.status == "RECEIVED"),
            "en_route": count(Complaint.status.in_(("ASSIGNED", "ACCEPTED", "EN_ROUTE"))),
            "on_site": count(Complaint.status.in_(("REACHED", "REPORT_RETURNED"))),
            "with_ro": count(Complaint.status == "REPORT_SUBMITTED"),
            "escalated_open": count(Complaint.status == "ESCALATED"),
            "sla_breached": count(Complaint.status.in_(OPEN), Complaint.is_sla_breached.is_(True)),
            "closed_today": count(Complaint.closed_at >= today)}


@router.get("/dashboard/sla-pipeline", summary="B2 Live count and average time per stage against target")
def pipeline(district: str | None = None, ac: str | None = None,
             date_from: datetime | None = Query(None, alias="from"), date_to: datetime | None = Query(None, alias="to"),
             user: User = Depends(dash_roles), db: Session = Depends(get_db)):
    conds = _scope_conditions(db, user, district, ac)
    start, end = _range(date_from, date_to)
    cfg = sla.load_config(db)
    base = select(ComplaintStageSla).join(Complaint, Complaint.id == ComplaintStageSla.complaint_id).where(*conds)
    out = []
    for stage in sla.STAGES:
        live = db.scalar(select(func.count()).select_from(base.where(
            ComplaintStageSla.stage == stage, ComplaintStageSla.ended_at.is_(None)).subquery()))
        breached_live = db.scalar(select(func.count()).select_from(base.where(
            ComplaintStageSla.stage == stage, ComplaintStageSla.ended_at.is_(None),
            ComplaintStageSla.deadline_at < utcnow()).subquery()))
        done = base.where(ComplaintStageSla.stage == stage, ComplaintStageSla.ended_at.between(start, end)).subquery()
        avg_sec, n, met = db.execute(select(func.avg(done.c.duration_sec), func.count(),
                                            func.sum(case((done.c.breached_at.is_(None), 1), else_=0)))).one()
        out.append({"stage": stage, "owner_role": sla.STAGE_OWNER[stage], "target_min": cfg[stage].target_min,
                    "live": live, "live_breached": breached_live,
                    "completed": n, "avg_sec": round(float(avg_sec)) if avg_sec is not None else None,
                    "within_target_pct": round(100 * int(met or 0) / n, 1) if n else None})
    return {"from": start, "to": end, "stages": out}


@router.get("/dashboard/outcomes", summary="B3 Disposed, dropped and escalated counts for a date range")
def outcomes(district: str | None = None, ac: str | None = None,
             date_from: datetime | None = Query(None, alias="from"), date_to: datetime | None = Query(None, alias="to"),
             user: User = Depends(dash_roles), db: Session = Depends(get_db)):
    conds = _scope_conditions(db, user, district, ac)
    start, end = _range(date_from, date_to)
    rows = dict(db.execute(select(Decision.action, func.count()).join(Complaint, Complaint.id == Decision.complaint_id)
                           .where(*conds, Decision.decided_at.between(start, end)).group_by(Decision.action)).all())
    dc = dict(db.execute(select(Complaint.status, func.count()).where(
        *conds, Complaint.status.in_(("DROPPED_AT_DC", "DUPLICATE", "RESOLVED", "WITHDRAWN")),
        Complaint.closed_at.between(start, end)).group_by(Complaint.status)).all())
    return {"from": start, "to": end,
            "disposed": rows.get("DISPOSE", 0), "dropped": rows.get("DROP", 0), "escalated": rows.get("ESCALATE", 0),
            "dropped_at_dc": dc.get("DROPPED_AT_DC", 0), "duplicates": dc.get("DUPLICATE", 0),
            "resolved_by_escalation": dc.get("RESOLVED", 0), "withdrawn": dc.get("WITHDRAWN", 0)}


GROUPS = {
    "district": (District.code, District.name_en),
    "ac": (Constituency.code, Constituency.name_en),
    "squad": (FlyingSquad.code, FlyingSquad.code),
    "violation_type": (ViolationType.code, ViolationType.name_en),
    "day": (func.date(Complaint.submitted_at), func.date(Complaint.submitted_at)),
}


@router.get("/reports/sla", summary="B4 SLA report grouped by district, AC, squad, violation type or day")
def sla_report(group_by: Literal["district", "ac", "squad", "violation_type", "day"] = "ac",
               district: str | None = None, ac: str | None = None,
               date_from: datetime | None = Query(None, alias="from"), date_to: datetime | None = Query(None, alias="to"),
               format: Literal["json", "csv", "xlsx"] = "json",
               user: User = Depends(require_roles("DC", "DEO", "CEO", "ADMIN")), db: Session = Depends(get_db)):
    conds = _scope_conditions(db, user, district, ac)
    start, end = _range(date_from, date_to)
    key, label = GROUPS[group_by]
    cur_dec = and_(Decision.complaint_id == Complaint.id, Decision.is_current.is_(True))
    q = (select(key.label("key"), label.label("label"),
                func.count(Complaint.id.distinct()).label("total"),
                func.sum(case((Complaint.status.in_(OPEN), 1), else_=0)).label("open"),
                func.sum(case((Complaint.status == "DISPOSED", 1), else_=0)).label("disposed"),
                func.sum(case((Complaint.status.in_(("DROPPED", "DROPPED_AT_DC")), 1), else_=0)).label("dropped"),
                func.sum(case((Complaint.status.in_(("ESCALATED", "RESOLVED")), 1), else_=0)).label("escalated"),
                func.sum(case((Complaint.status == "DUPLICATE", 1), else_=0)).label("duplicate"),
                func.sum(case((Complaint.is_sla_breached.is_(True), 1), else_=0)).label("stage_breaches"),
                func.sum(case((Decision.sla_met.is_(True), 1), else_=0)).label("within_100_min"),
                func.avg(Decision.total_minutes).label("avg_total_min"))
         .select_from(Complaint)
         .outerjoin(Decision, cur_dec)
         .join(District, District.id == Complaint.district_id)
         .outerjoin(Constituency, Constituency.id == Complaint.ac_id)
         .join(ViolationType, ViolationType.id == Complaint.violation_type_id)
         .outerjoin(Assignment, and_(Assignment.complaint_id == Complaint.id, Assignment.is_active.is_(True)))
         .outerjoin(FlyingSquad, FlyingSquad.id == Assignment.squad_id)
         .where(*conds, Complaint.submitted_at.between(start, end))
         .group_by(key, label).order_by(key))
    rows = []
    for r in db.execute(q):
        decided = int(r.disposed or 0) + int(r.dropped or 0) + int(r.escalated or 0)
        rows.append({"key": str(r.key) if r.key is not None else "(none)", "label": str(r.label or "(none)"),
                     "total": r.total, "open": int(r.open or 0), "disposed": int(r.disposed or 0),
                     "dropped": int(r.dropped or 0), "escalated": int(r.escalated or 0),
                     "duplicate": int(r.duplicate or 0), "stage_breaches": int(r.stage_breaches or 0),
                     "within_100_min": int(r.within_100_min or 0),
                     "within_100_min_pct": round(100 * int(r.within_100_min or 0) / decided, 1) if decided else None,
                     "avg_total_min": round(float(r.avg_total_min), 1) if r.avg_total_min is not None else None})
    if format == "json":
        return {"group_by": group_by, "from": start, "to": end, "rows": rows}
    cols = list(rows[0].keys()) if rows else ["key", "label", "total"]
    fname = f"sla_report_{group_by}_{start:%Y%m%d}_{end:%Y%m%d}"
    if format == "csv":
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)
        return Response(buf.getvalue().encode("utf-8-sig"), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{fname}.csv"'})
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "SLA report"
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r[c] for c in cols])
    ws.freeze_panes = "A2"
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, len(c) + 2)
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'})


@router.get("/complaints/{complaint_id}/audit-log", summary="B5 Every action on a complaint: who, what, when")
def audit_log(complaint_id: str, user: User = Depends(require_roles(*OFFICIAL_ROLES)), db: Session = Depends(get_db)):
    if user.role == "FS":
        raise forbidden("Audit log is not available to flying squads")
    c = get_complaint(db, complaint_id)
    assert_can_view(db, user, c)
    rows = db.execute(select(AuditLog, User.name, User.designation)
                      .outerjoin(User, User.id == AuditLog.actor_id)
                      .where(AuditLog.entity == "complaint", AuditLog.entity_id == c.id)
                      .order_by(AuditLog.created_at, AuditLog.id)).all()
    return {"complaint_no": c.complaint_no,
            "results": [{"id": a.id, "action": a.action, "actor_role": a.actor_role,
                         "actor": ({"name": name, "designation": desig} if a.actor_role != "CITIZEN" else None),
                         "ip_address": a.ip_address, "user_agent": a.user_agent, "payload": a.payload,
                         "at": a.created_at} for a, name, desig in rows]}
