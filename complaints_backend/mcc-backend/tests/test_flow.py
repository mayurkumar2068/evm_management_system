"""End-to-end tests against a real MySQL 8 database."""
import io
from datetime import timedelta

import pytest
from openpyxl import load_workbook

from tests.conftest import API, MADHYA, NARELA, OUTSIDE, jpeg_bytes, now_iso


def test_api_count_is_55(client):
    from app.main import app
    ops = sum(len(v) for v in app.openapi()["paths"].values())
    assert ops == 54                     # REST
    assert client.app is app
    with pytest.raises(Exception):                                   # + 1 WebSocket (rejects bad token)
        with client.websocket_connect(f"{API}/ws/live?token=bad") as ws:
            ws.receive_text()


def test_standard_error_format(api):
    r = api.c.get(f"{API}/users/me")
    assert r.status_code == 401
    assert r.json()["error"]["code"] == "UNAUTHORIZED"
    r = api.post("/complaints", api.citizen("9800000099"), json={"violation_type": "X"})
    assert r.status_code == 400 and r.json()["error"]["code"] == "VALIDATION_ERROR"


def test_full_happy_path(api):
    cit = api.citizen("9800000001")
    vt = api.get("/violation-types", cit).json()["results"]
    assert any(v["code"] == "LIQUOR_DISTRIBUTION" and v["name_hi"] for v in vt)

    c = api.new_complaint("9800000001")
    assert c["status"] == "RECEIVED"
    assert c["complaint_no"].startswith("BPL-") and c["ac"]["code"] == "AC-152"
    cid = c["id"]

    dc = api.official("dc.bhopal")
    q = api.get("/dc/complaints?status=unassigned", dc).json()
    assert any(x["id"] == cid for x in q["results"])
    row = next(x for x in q["results"] if x["id"] == cid)
    assert row["current_stage"] == "ALLOCATION" and 0 < row["seconds_left"] <= 300

    near = api.get(f"/dc/complaints/{cid}/nearby-squads", dc).json()["results"]
    assert near[0]["code"] == "FS-02" and near[0]["distance_km"] < 2

    detail = api.get(f"/complaints/{cid}", dc).json()
    assert detail["citizen"] is None                       # anonymous
    assert detail["media"][0]["url"].startswith("http")

    r = api.post(f"/dc/complaints/{cid}/assign", dc, json={"squad_id": near[0]["squad_id"], "note": "Go fast"})
    assert r.status_code == 200 and r.json()["status"] == "ASSIGNED"
    aid = r.json()["assignment_id"]

    fs = api.official("fs02.lead")
    mine = api.get("/fs/assignments", fs).json()["results"]
    assert any(a["id"] == aid for a in mine)
    assert api.post(f"/fs/assignments/{aid}/accept", fs).json()["status"] == "ACCEPTED"
    start = api.post(f"/fs/assignments/{aid}/start", fs).json()
    assert start["status"] == "EN_ROUTE" and "google.com/maps" in start["navigation_url"]
    r = api.post("/fs/location", fs, json={"points": [{"lat": 23.262, "lng": 77.431, "at": now_iso()}]})
    assert r.json()["accepted"] == 1

    far = api.post(f"/fs/assignments/{aid}/reached", fs, json={"lat": 23.2601, "lng": 77.4290})
    assert far.status_code == 422 and far.json()["error"]["code"] == "GEO_MISMATCH"
    ok = api.post(f"/fs/assignments/{aid}/reached", fs, json={"lat": NARELA[0] + 0.0004, "lng": NARELA[1]})
    assert ok.status_code == 200 and ok.json()["distance_m"] < 200

    pdf = b"%PDF-1.4\n% demo panchnama\n"
    r = api.post(f"/fs/assignments/{aid}/media", fs, files={"file": ("p.pdf", pdf, "application/pdf")},
                 data={"media_type": "DOCUMENT", "doc_kind": "PANCHNAMA", "captured_at": now_iso()})
    assert r.status_code == 201, r.text
    media_id = r.json()["id"]
    r = api.post(f"/fs/assignments/{aid}/report", fs, json={
        "finding": "VIOLATION_CONFIRMED", "summary": "48 bottles of liquor found in a parked van",
        "party_or_candidate": "[recorded]", "actions_taken": ["SEIZURE", "PANCHNAMA"], "media_ids": [media_id],
        "seizure": {"made": True, "items": [{"type": "LIQUOR", "quantity": 48, "unit": "bottle", "value_inr": 19200}]},
        "recommendation": "DISPOSE"})
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "REPORT_SUBMITTED"

    ro = api.official("ro.narela")
    assert any(x["id"] == cid for x in api.get("/ro/complaints", ro).json()["results"])
    rep = api.get(f"/ro/complaints/{cid}/field-report", ro).json()
    assert rep["report"]["seizure"]["items"][0]["value_inr"] == 19200
    assert len(rep["media"]["squad"]) == 1 and len(rep["media"]["citizen"]) == 1
    assert [s["stage"] for s in rep["stages"]] == ["ALLOCATION", "TRAVEL", "ENQUIRY", "RO_DECISION"]

    bad = api.post(f"/ro/complaints/{cid}/decision", ro, json={"action": "DISPOSE", "remarks": "Action taken now"})
    assert bad.status_code == 400
    r = api.post(f"/ro/complaints/{cid}/decision", ro, json={
        "action": "DISPOSE", "final_action": "FIR_REGISTERED", "fir_no": "123/2026",
        "remarks": "FIR registered and liquor seized"})
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "DISPOSED" and r.json()["sla_met"] is True

    view = api.get(f"/citizen/complaints/{cid}", cit).json()
    assert view["status"] == "DISPOSED" and view["remarks"] == "FIR registered and liquor seized"
    assert [t["status"] for t in view["timeline"]] == ["RECEIVED", "ASSIGNED", "REACHED", "REPORT_SUBMITTED",
                                                        "DISPOSED"]
    assert "decided_by" not in str(view) and "fs02" not in str(view)
    notes = api.get("/notifications", cit).json()
    assert notes["unread"] >= 5
    assert api.post(f"/notifications/{notes['results'][0]['id']}/read", cit).json()["read"] is True

    assert any(d["complaint_id"] == cid for d in api.get("/ro/decisions", ro).json()["results"])
    audit = api.get(f"/complaints/{cid}/audit-log", dc).json()["results"]
    assert [a["action"] for a in audit][-1] == "STATUS_DISPOSED"
    assert api.get(f"/complaints/{cid}/audit-log", fs).status_code == 403

    # media link works and is signed
    url = detail["media"][0]["url"].replace("http://localhost:8000", "")
    assert api.c.get(url).status_code == 200
    assert api.c.get(url.split("&sig=")[0] + "&sig=bad").status_code == 403


def test_dashboard_and_reports(api):
    api.to_report_submitted(mobile="9800000002")
    dc = api.official("dc.bhopal")
    s = api.get("/dashboard/summary", dc).json()
    assert s["received_today"] >= 2 and s["with_ro"] >= 1
    p = api.get("/dashboard/sla-pipeline", dc).json()
    stages = {x["stage"]: x for x in p["stages"]}
    assert stages["RO_DECISION"]["live"] >= 1 and stages["ALLOCATION"]["completed"] >= 2
    assert stages["ALLOCATION"]["target_min"] == 5
    o = api.get("/dashboard/outcomes", dc).json()
    assert o["disposed"] >= 1
    rep = api.get("/reports/sla?group_by=violation_type", dc).json()
    assert rep["rows"] and rep["rows"][0]["total"] >= 1
    csv = api.get("/reports/sla?group_by=ac&format=csv", dc)
    assert csv.headers["content-type"].startswith("text/csv") and b"within_100_min" in csv.content
    xl = api.get("/reports/sla?group_by=squad&format=xlsx", dc)
    ws = load_workbook(io.BytesIO(xl.content)).active
    assert ws["A1"].value == "key" and ws.max_row >= 2
    ro = api.official("ro.narela")
    assert api.get("/dashboard/summary", ro).status_code == 200
    assert api.get("/reports/sla", ro).status_code == 403


def test_scope_isolation(api):
    c = api.new_complaint("9800000003")
    ro_other = api.official("ro.madhya")
    assert api.get(f"/complaints/{c['id']}", ro_other).status_code == 404
    other_cit = api.citizen("9800000004")
    assert api.get(f"/citizen/complaints/{c['id']}", other_cit).status_code == 404
    assert api.get("/dc/complaints", other_cit).status_code == 403
    dc = api.official("dc.bhopal")
    a = api.post(f"/dc/complaints/{c['id']}/assign", dc, json={"squad_id": api.squad_id("FS-01")}).json()
    fs3 = api.official("fs03.lead")
    assert api.post(f"/fs/assignments/{a['assignment_id']}/accept", fs3).status_code == 404
    assert api.get(f"/complaints/{c['id']}", fs3).status_code == 404
    assert api.get(f"/complaints/{c['id']}", api.official("fs01.lead")).status_code == 200


def test_invalid_state_and_rules(api):
    c = api.new_complaint("9800000005")
    dc, ro = api.official("dc.bhopal"), api.official("ro.narela")
    r = api.post(f"/ro/complaints/{c['id']}/decision", ro, json={
        "action": "DROP", "reason_code": "NOTHING_FOUND", "remarks": "Nothing found here"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "INVALID_STATE"
    sid = api.squad_id("FS-03")
    assert api.post(f"/dc/complaints/{c['id']}/assign", dc, json={"squad_id": sid}).status_code == 200
    again = api.post(f"/dc/complaints/{c['id']}/assign", dc, json={"squad_id": sid})
    assert again.status_code == 409
    w = api.post(f"/complaints/{c['id']}/withdraw", c["h"])
    assert w.status_code == 409
    # busy squad cannot be assigned
    c2 = api.new_complaint("9800000006")
    assert api.post(f"/dc/complaints/{c2['id']}/assign", dc, json={"squad_id": sid}).status_code == 409
    # squad cannot go off duty during an assignment
    assert api.patch("/fs/me/status", api.official("fs03.lead"), json={"status": "OFF_DUTY"}).status_code == 409
    # reassign frees FS-03
    r = api.post(f"/dc/complaints/{c['id']}/reassign", dc,
                 json={"squad_id": api.squad_id("FS-02"), "reason": "FS-03 vehicle breakdown"})
    assert r.status_code == 200, r.text
    assert api.patch("/fs/me/status", api.official("fs03.lead"), json={"status": "AVAILABLE"}).status_code == 200
    detail = api.get(f"/complaints/{c['id']}", dc).json()
    assert [a["is_active"] for a in detail["assignments"]] == [False, True]
    # citizen rules
    cit = api.citizen("9800000007")
    old = (__import__("datetime").datetime.now(__import__("datetime").timezone.utc) - timedelta(minutes=10)).isoformat()
    r = api.post("/complaints", cit, json={"violation_type": "LIQUOR_DISTRIBUTION", "latitude": NARELA[0],
                                           "longitude": NARELA[1], "captured_at": old})
    assert r.status_code == 400
    d = api.new_complaint("9800000007", submit=False)
    r = api.post(f"/complaints/{d['id']}/media", cit, files={"file": ("x.txt", b"hello", "text/plain")},
                 data={"media_type": "PHOTO", "captured_at": now_iso()})
    assert r.status_code == 400
    out = api.new_complaint("9800000008", where=OUTSIDE, submit=False)
    r = api.post(f"/complaints/{out['id']}/submit", out["h"])
    assert r.status_code == 422 and r.json()["error"]["code"] == "GEO_MISMATCH"
    assert api.post(f"/complaints/{out['id']}/withdraw", out["h"]).json()["status"] == "WITHDRAWN"


def test_duplicate_and_drop(api):
    a = api.new_complaint("9800000010", where=MADHYA, vtype="MONEY_DISTRIBUTION")
    assert a["ac"]["code"] == "AC-151"
    b = api.new_complaint("9800000011", where=(MADHYA[0] + 0.0005, MADHYA[1]), vtype="MONEY_DISTRIBUTION")
    assert b["duplicate_of"] == a["complaint_no"]
    dc = api.official("dc.bhopal")
    r = api.post(f"/dc/complaints/{b['id']}/mark-duplicate", dc, json={"original_complaint": a["complaint_no"]})
    assert r.status_code == 200 and r.json()["status"] == "DUPLICATE"
    r = api.post(f"/dc/complaints/{a['id']}/drop", dc, json={"reason_code": "INSUFFICIENT_DETAILS",
                                                               "remarks": "Photo does not show any activity"})
    assert r.json()["status"] == "DROPPED_AT_DC"
    v = api.get(f"/citizen/complaints/{a['id']}", a["h"]).json()
    assert v["remarks"] == "Photo does not show any activity"


def test_send_back_then_escalation_chain(api):
    c = api.to_report_submitted(mobile="9800000012", squad="FS-03", fs_user="fs03.lead")
    ro = api.official("ro.narela")
    r = api.post(f"/ro/complaints/{c['id']}/send-back", ro, json={"remarks": "Attach the seizure memo please"})
    assert r.json()["status"] == "REPORT_RETURNED"
    fs = api.official("fs03.lead")
    r = api.post(f"/fs/assignments/{c['assignment_id']}/report", fs, json={
        "finding": "VIOLATION_CONFIRMED", "summary": "Seizure memo attached, 48 bottles"})
    assert r.json()["version"] == 2
    r = api.post(f"/ro/complaints/{c['id']}/decision", ro, json={
        "action": "ESCALATE", "escalate_to": "DEO", "remarks": "Involves a sitting minister"})
    assert r.json()["status"] == "ESCALATED"

    deo = api.official("deo.bhopal")
    lst = api.get("/escalations", deo).json()["results"]
    esc = next(e for e in lst if e["complaint"]["id"] == c["id"])
    full = api.get(f"/escalations/{esc['id']}", deo).json()
    assert full["complaint"]["field_reports"][0]["version"] == 2

    r = api.post(f"/escalations/{esc['id']}/action", deo, json={"action": "FORWARD", "remarks": "Needs CEO review"})
    assert r.json()["new_escalation"]["level"] == "CEO"
    ceo = api.official("ceo.mp")
    new_id = r.json()["new_escalation"]["id"]
    assert api.post(f"/escalations/{esc['id']}/action", deo,
                    json={"action": "RESOLVE", "final_action": "X", "remarks": "Already closed one"}).status_code == 409
    r = api.post(f"/escalations/{new_id}/action", ceo, json={"action": "RESOLVE", "final_action": "NOTICE_ISSUED",
                                                              "remarks": "Notice issued to the candidate"})
    assert r.json()["complaint_status"] == "RESOLVED"
    v = api.get(f"/citizen/complaints/{c['id']}", api.citizen("9800000012")).json()
    assert v["status"] == "RESOLVED" and v["remarks"] == "Notice issued to the candidate"


def test_escalation_return_to_ro(api):
    c = api.to_report_submitted(mobile="9800000013", squad="FS-03", fs_user="fs03.lead")
    ro, deo = api.official("ro.narela"), api.official("deo.bhopal")
    api.post(f"/ro/complaints/{c['id']}/decision", ro, json={
        "action": "ESCALATE", "escalate_to": "DEO", "remarks": "Not sure about jurisdiction"})
    esc = next(e for e in api.get("/escalations", deo).json()["results"] if e["complaint"]["id"] == c["id"])
    r = api.post(f"/escalations/{esc['id']}/action", deo, json={"action": "RETURN_TO_RO",
                                                                "remarks": "Within RO powers, decide"})
    assert r.json()["complaint_status"] == "REPORT_SUBMITTED"
    r = api.post(f"/ro/complaints/{c['id']}/decision", ro, json={
        "action": "DROP", "reason_code": "NOT_MCC", "remarks": "Private event, not MCC"})
    assert r.status_code == 200
    d = api.get(f"/complaints/{c['id']}", ro).json()["decisions"]
    assert [x["is_current"] for x in d] == [False, True]


def test_sla_worker_warning_and_breach(api):
    from sqlalchemy import update

    from app.db.session import SessionLocal
    from app.models import ComplaintStageSla, Complaint
    from app.services.sla import run_sla_check
    from app.core.timeutil import utcnow

    c = api.new_complaint("9800000014")
    db = SessionLocal()
    cid = db.query(Complaint.id).filter(Complaint.uuid == c["id"]).scalar()
    db.execute(update(ComplaintStageSla).where(ComplaintStageSla.complaint_id == cid).values(
        started_at=utcnow() - timedelta(minutes=4, seconds=30), deadline_at=utcnow() + timedelta(seconds=30)))
    db.commit()
    assert run_sla_check(db)["warned"] >= 1
    db.execute(update(ComplaintStageSla).where(ComplaintStageSla.complaint_id == cid).values(
        deadline_at=utcnow() - timedelta(seconds=1)))
    db.commit()
    assert run_sla_check(db)["breached"] >= 1
    assert run_sla_check(db)["breached"] == 0          # idempotent
    db.close()
    dc = api.official("dc.bhopal")
    q = api.get("/dc/complaints?sla_breached=true", dc).json()["results"]
    assert any(x["id"] == c["id"] for x in q)
    notes = api.get("/notifications", dc).json()["results"]
    assert any(n["type"] == "SLA_BREACHED" for n in notes)
    deo_notes = api.get("/notifications", api.official("deo.bhopal")).json()["results"]
    assert not any(n["type"] == "SLA_BREACHED" and n["data"].get("complaint_id") == c["id"] for n in deo_notes)


def test_websocket_live_events(api, client):
    import threading
    from app.core.deps import user_from_token  # noqa: F401
    dc = api.official("dc.bhopal")
    token = dc["Authorization"].split()[1]
    with client.websocket_connect(f"{API}/ws/live?token={token}") as ws:
        assert ws.receive_json()["event"] == "connected"
        box = {}
        t = threading.Thread(target=lambda: box.setdefault("c", api.new_complaint("9800000015")))
        t.start()
        t.join()
        events = [ws.receive_json()["event"] for _ in range(2)]
        assert "complaint.created" in events and "complaint.status_changed" in events
    ro2 = api.official("ro.madhya")
    with client.websocket_connect(f"{API}/ws/live?token={ro2['Authorization'].split()[1]}") as ws:
        ws.receive_json()
        api.new_complaint("9800000016")      # Narela: RO Madhya must not receive it
        ws.send_text("ping")
        assert ws.receive_text() == "pong"


def test_auth_tokens(api):
    c = api.c
    r = c.post(f"{API}/auth/otp/send", json={"mobile": "9800000020"}).json()
    bad = c.post(f"{API}/auth/otp/verify", json={"mobile": "9800000020", "otp_ref": r["otp_ref"],
                                                 "otp": "000000" if r["debug_otp"] != "000000" else "111111"})
    assert bad.status_code == 401
    ok = c.post(f"{API}/auth/otp/verify", json={"mobile": "9800000020", "otp_ref": r["otp_ref"],
                                                "otp": r["debug_otp"]}).json()
    reuse = c.post(f"{API}/auth/otp/verify", json={"mobile": "9800000020", "otp_ref": r["otp_ref"],
                                                   "otp": r["debug_otp"]})
    assert reuse.status_code == 401
    new = c.post(f"{API}/auth/token/refresh", json={"refresh": ok["refresh"]})
    assert new.status_code == 200
    assert c.post(f"{API}/auth/token/refresh", json={"refresh": ok["refresh"]}).status_code == 401
    h = {"Authorization": f"Bearer {new.json()['access']}"}
    assert c.get(f"{API}/users/me", headers=h).json()["role"] == "CITIZEN"
    assert c.post(f"{API}/devices", headers=h, json={"platform": "ANDROID", "fcm_token": "tok-abcdefghijk"}).status_code == 201
    assert c.post(f"{API}/auth/logout", headers=h, json={"refresh": new.json()["refresh"],
                                                         "fcm_token": "tok-abcdefghijk"}).json()["ok"]
    assert c.post(f"{API}/auth/token/refresh", json={"refresh": new.json()["refresh"]}).status_code == 401
    assert c.post(f"{API}/auth/login", json={"username": "dc.bhopal", "password": "wrong-pass"}).status_code == 401


def test_admin_apis(api):
    adm = api.official("admin")
    r = api.post("/admin/squads", adm, json={"code": "FS-09", "district_code": "BPL", "vehicle_no": "MP04-9900",
                                             "area_ac_codes": ["AC-152"]})
    assert r.status_code == 201, r.text
    sq = r.json()
    r = api.post("/admin/users", adm, json={"role": "FS", "name": "New Squad Member", "mobile": "9111111111",
                                            "username": "fs09.member", "password": "Strong@123",
                                            "district_code": "BPL", "squad_id": sq["id"]})
    assert r.status_code == 201, r.text
    uid = r.json()["id"]
    assert api.post("/admin/users", adm, json={"role": "RO", "name": "No AC", "mobile": "9111111112",
                                               "username": "ro.noac", "password": "Strong@123",
                                               "district_code": "BPL"}).status_code == 400
    r = api.patch(f"/admin/squads/{sq['id']}", adm, json={"member_user_ids": [uid], "leader_user_id": uid,
                                                          "status": "AVAILABLE"})
    assert r.json()["leader"]["id"] == uid and len(r.json()["members"]) == 1
    assert api.patch(f"/admin/users/{uid}", adm, json={"designation": "Naib Tehsildar"}).status_code == 200
    assert any(u["username"] == "fs09.member" for u in api.get("/admin/users?role=FS", adm).json()["results"])
    assert any(s["code"] == "FS-09" for s in api.get("/admin/squads", adm).json()["results"])
    ds = api.get("/masters/districts?include_boundary=true", adm).json()["results"]
    assert ds[0]["boundary"]["type"] == "MultiPolygon"
    acs = api.get("/masters/constituencies?district=BPL", adm).json()["results"]
    assert {a["code"] for a in acs} == {"AC-151", "AC-152"} and all(a["ro"] for a in acs)
    r = api.put("/admin/sla-config", adm, json={"stages": [{"stage": "ALLOCATION", "target_min": 5, "warn_pct": 70,
                                                            "escalate_to_role": "DC"}]})
    assert r.json()["total_min"] == 100
    assert api.get("/admin/users", api.official("dc.bhopal")).status_code == 403
    assert api.patch(f"/admin/users/{uid}", adm, json={"is_active": False}).status_code == 200
